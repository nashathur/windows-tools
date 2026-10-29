"""Phase 6 document conversions: office2pdf, toword, md2pdf (+ check --source), txt2pdf."""

from __future__ import annotations

import contextlib
import difflib
import html
import html.entities as html_entities
import json
import os
import re
import shutil
import sys
import unicodedata
from pathlib import Path

import pymupdf

import clierr
from .. import crypt, io, tools
from ..errors import PdfError
from ..pipeline import open_input, _signed
from ..registry import command

ROOT = Path(__file__).resolve().parents[2]
NODE_MODULES = ROOT / "node_modules"


def _path(value):
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise PdfError("NOT_FOUND", f"input not found: {path}", "check the path")
    return path


def _result(command_name, output, pages=None, warnings=(), verdict=None, kind="pdf"):
    warnings = list(warnings)
    return {"type": "result", "command": command_name, "path": str(output), "kind": kind,
            "pages": pages, "bytes": output.stat().st_size,
            "verdict": verdict or ("warn" if warnings else "report"), "warnings": warnings}


def _dry_run(command_name, output, kind="pdf", pages=None):
    return {"type": "result", "command": command_name, "path": str(output), "kind": kind,
            "pages": pages, "bytes": 0, "verdict": "report", "provisional": True, "warnings": []}


def _fail_tool(exc):
    tool_run = getattr(exc, "tool_run", None)
    detail = getattr(exc, "detail", None) or (
        f"argv: {getattr(tool_run, 'args', None)}\nstdout:\n{getattr(tool_run, 'stdout', '')}\nstderr:\n{getattr(tool_run, 'stderr', '')}"
        if tool_run else None)
    return PdfError(exc.code, exc.what, exc.fix, detail=detail)


def _unchecked(command_name, stage, output, context, problems, detail=None):
    """Publish under the _UNCHECKED name, exit 3, one WARN per problem, one log for the measurements."""
    stage.target = output.with_name(output.stem + "_UNCHECKED" + output.suffix)
    context.exit_code = 3
    log = clierr.write_log("pdf", "CHECK_FAILED", f"{command_name} invariant check failed",
                           "review the converted document against its source",
                           sections={"findings": json.dumps([{"code": c, "message": m} for c, m in problems],
                                                            ensure_ascii=False),
                                     **({"detail": str(detail)} if detail else {})},
                           log_name=command_name)
    for code, message in problems:
        clierr.warn(f"pdf {command_name}", code, message, "review the converted document")
    return log


def _publish_pdf(source, args, context, writer, *, command_name, verify=None):
    """Convert `source` with writer(source, staged_path, temp) and publish the PDF.

    verify(staged_pdf) -> [(CODE, message)] is the command's real check; without one the verdict is
    `report` (nothing was verified). A problem publishes `<name>_UNCHECKED.pdf`, exit 3."""
    source = _path(source)
    output = io.resolve_output(source, command_name, getattr(args, "output", None), suffix="", extension=".pdf")
    if args.dry_run:
        return _dry_run(command_name, output)
    output.parent.mkdir(parents=True, exist_ok=True)
    io.sweep_stale(output.parent)
    problems, log = [], None
    with io.temp_folder() as temp:
        io.disk_budget([source], output, temp, multiplier=3)
        with io.Staging(output, temp) as stage:
            writer(source, stage.path, Path(temp))
            try:
                with pymupdf.open(stage.path) as doc:
                    if not doc.page_count:
                        raise ValueError("converted PDF contains no pages")
                    pages = doc.page_count
            except Exception as exc:
                raise PdfError("TOOL_FAILED", f"conversion did not produce a readable PDF: {exc}",
                               "check the source document and installed converter", detail=str(exc))
            verification_warnings = []
            if verify:
                verified = verify(stage.path)
                if isinstance(verified, dict):
                    problems = verified.get("problems", [])
                    verification_warnings = verified.get("warnings", [])
                else:
                    problems = verified
                if problems:
                    log = _unchecked(command_name, stage, output, context, problems)
            final = stage.publish()
    warnings = [c for c, _ in problems] + [c for c, _ in verification_warnings]
    for code, message in verification_warnings:
        clierr.warn(f"pdf {command_name}", code, message, "review the converted document")
    if final != stage.target:
        warnings.append("NAME_TAKEN")
        clierr.warn(f"pdf {command_name}", "NAME_TAKEN", f"published as {final.name}", "choose another output name")
    record = _result(command_name, final, pages, warnings,
                     verdict="unchecked" if problems else ("warn" if warnings else ("ok" if verify else "report")))
    if log:
        record["log"] = log
    return record


# ---------------------------------------------------------------------------------------------
# Office -> PDF

def add_arguments_office2pdf(parser):
    parser.add_argument("file", type=Path)
    parser.add_argument("-o", "--output")


# LibreOffice 26.8 schema (share\registry\main.xcd). Update/Link: Writer 0 never, 1 on request,
# 2 always; Calc 0 always, 1 never, 2 on request. Macro level 3 = very high; documents binding
# macros to events are blocked (CheckDocumentEvents) and links from untrusted places are refused.
OFFICE_SETTINGS = (
    ("/org.openoffice.Office.Common/Security/Scripting",
     (("MacroSecurityLevel", "3"), ("DisableMacrosExecution", "true"),
      ("CheckDocumentEvents", "true"), ("BlockUntrustedRefererLinks", "true"))),
    ("/org.openoffice.Office.Writer/Content/Update", (("Link", "0"),)),
    ("/org.openoffice.Office.Calc/Content/Update", (("Link", "1"),)),
)


def _office_security(profile):
    user = Path(profile) / "user"
    user.mkdir(parents=True, exist_ok=True)
    items = "".join(
        f'<item oor:path="{path}">' + "".join(
            f'<prop oor:name="{name}" oor:op="fuse"><value>{value}</value></prop>' for name, value in props)
        + "</item>\n" for path, props in OFFICE_SETTINGS)
    (user / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
        'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
        + items + "</oor:items>\n", encoding="utf-8")


def office2pdf(args, context):
    source = _path(args.file)
    soffice = tools.find_tool("soffice")
    if not soffice:
        raise PdfError("NOT_FOUND", "LibreOffice is not installed", "install LibreOffice")

    def writer(source, staged, temp):
        profile = temp / "lo-profile"
        _office_security(profile)
        outdir = temp / "lo-output"
        outdir.mkdir()
        argv = [soffice, "--headless", "--norestore", "--nolockcheck",
                f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf",
                "--outdir", str(outdir), str(source)]
        try:
            result = tools.run_tool(argv, timeout=120, cwd=temp, temp=temp,
                                    secret_env=getattr(context, "secret_env", ()))
        except clierr.Fail as exc:
            raise _fail_tool(exc)
        made = outdir / (source.stem + ".pdf")
        if not made.is_file():
            raise PdfError("TOOL_FAILED", "LibreOffice exited without creating a PDF",
                           "check the source document and LibreOffice log", detail=result.stderr)
        shutil.move(made, staged)
    # No invariant is checked (the source has no portable text count): verdict `report`.
    return _publish_pdf(source, args, context, writer, command_name="office2pdf")


# ---------------------------------------------------------------------------------------------
# PDF -> Word

def add_arguments_toword(parser):
    parser.add_argument("file", type=Path)
    parser.add_argument("-o", "--output")
    parser.add_argument("--password-file", type=Path)
    parser.add_argument("--password-env")


# Runs in a child interpreter so a hung conversion can be killed. The password arrives in a file in
# the run's temp folder, never in argv.
TOWORD_CHILD = """\
import json, sys
from pdf2docx import Converter
source, output, password = json.load(open(sys.argv[1], encoding='utf-8'))
cv = Converter(source, password=password)
try:
    cv.convert(output, ignore_page_error=False)
finally:
    cv.close()
"""


def _failed_page_note(text):
    """' on page N' from pdf2docx's error line; its `(1/7) Page 1` progress lines never name a failure."""
    errors = re.findall(r"(?:Error when parsing page|error on page|failed to parse page)\s*(\d+)", text, re.I)
    return f" on page {errors[-1]}" if errors else ""


def _docx_words(document):
    """Words in paragraphs and table cells; a merged cell repeats in row.cells but counts once."""
    count = sum(len(p.text.split()) for p in document.paragraphs)
    seen = set()
    seen_cells = []  # retain XML elements so their ids cannot be reused
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                identity = id(cell._tc)
                if identity not in seen:
                    seen.add(identity)
                    seen_cells.append(cell._tc)
                    count += len(cell.text.split())
    return count


def toword(args, context):
    source = _path(args.file)
    password = context.password or context.owner_password
    with open_input(source, password) as doc:
        pages = doc.page_count
        pdf_words = sum(len(page.get_text("words")) for page in doc)
        encrypted = bool(doc.is_encrypted or doc.metadata.get("encryption"))  # is_encrypted turns False once authenticated
        signed = _signed(source, password)
    output = io.resolve_output(source, "toword", args.output, suffix="", extension=".docx")
    warnings = crypt.export_warnings(encrypted)
    if signed:
        warnings.append("SIGNATURE_NOT_RETAINED")
    if args.dry_run:
        record = _dry_run("toword", output, "docx", pages)
        record["warnings"] = warnings
        return record
    output.parent.mkdir(parents=True, exist_ok=True)
    io.sweep_stale(output.parent)
    with io.temp_folder() as temp, io.Staging(output, temp) as stage:
        io.disk_budget([source], output, temp, multiplier=3)
        job, script = Path(temp) / "toword.json", Path(temp) / "toword_child.py"
        job.write_text(json.dumps([str(source), str(stage.path), password]), encoding="utf-8")
        script.write_text(TOWORD_CHILD, encoding="utf-8")
        store = getattr(context, "secret_store", None)
        try:
            tools.run_tool([sys.executable, "-I", str(script), str(job)], cwd=temp,
                           timeout=max(120, 20 * pages), temp=temp, secret_env=context.secret_env,
                           secret_values=tuple(store.values) if store else ())
        except clierr.Fail as exc:
            stderr = getattr(getattr(exc, "tool_run", None), "stderr", None) or ""
            detail = getattr(exc, "detail", None) or stderr or exc.what
            where = _failed_page_note(stderr + "\n" + detail)
            raise PdfError(exc.code, f"pdf2docx conversion failed{where}: {exc.what}",
                           "check the PDF and pdf2docx log", detail=detail)
        finally:
            job.unlink(missing_ok=True)
        from docx import Document
        try:
            document = Document(stage.path)
            word_count = _docx_words(document)
        except Exception as exc:
            raise PdfError("TOOL_FAILED", "pdf2docx produced an invalid DOCX", "retry with a readable PDF", detail=str(exc))
        if pdf_words and word_count < pdf_words * 0.8:
            warnings.append("LOW_TEXT")
        final = stage.publish()
    if final != stage.target:
        warnings.append("NAME_TAKEN")
    for code in warnings:
        detail = ("toword exported an encrypted PDF to an unencrypted DOCX" if code == "DECRYPTED_EXPORT"
                  else "the DOCX does not retain the PDF signature" if code == "SIGNATURE_NOT_RETAINED"
                  else f"toword produced warning {code}")
        clierr.warn("pdf toword", code, detail, "review the exported document")
    # The word ratio is the check; a PDF without text (scan) gives nothing to compare: `report`.
    return _result("toword", final, pages, warnings, kind="docx",
                   verdict="warn" if warnings else ("ok" if pdf_words else "report"))


# ---------------------------------------------------------------------------------------------
# Markdown -> PDF and check --source

def add_arguments_md2pdf(parser):
    parser.add_argument("file", type=Path)
    parser.add_argument("-o", "--output")


def _node_env(**extra):
    return {**os.environ, "NODE_PATH": str(NODE_MODULES), **extra}


def _markdown_diagnostics(warnings, source_warnings, problems):
    shown = {code for code, _ in problems}
    source_codes = {warning.get("code", "SOURCE_WARNING") for warning in source_warnings}
    diagnostics, seen = [], set()
    for warning in source_warnings:
        code = warning.get("code", "SOURCE_WARNING")
        detail = ", ".join(f"{key} {warning[key]}" for key in ("page", "section") if warning.get(key))
        key = (code, detail)
        if code not in shown and key not in seen:
            diagnostics.append((code, f"Markdown conversion warning{': ' + detail if detail else ''}"))
            seen.add(key)
    for code in dict.fromkeys(warnings):
        if code not in shown and code not in source_codes and code not in seen:
            diagnostics.append((code, "Markdown conversion completed with a warning"))
            seen.add(code)
    return diagnostics


def _clean_markdown_path(path):
    path = _path(path)
    if str(path).startswith("\\\\"):
        raise PdfError("BAD_ARG", "UNC Markdown paths are not allowed", "copy the Markdown file to a local folder")
    return path


def md2pdf(args, context):
    source = _clean_markdown_path(args.file)
    node = shutil.which("node")
    edge = tools.find_tool("browser")
    if not node or not edge:
        missing = "Node.js" if not node else "Microsoft Edge or Chrome"
        raise PdfError("NOT_FOUND", f"{missing} is not installed", "install the required renderer")
    script = ROOT / "md2pdf.js"
    # Absolute: node runs in the temp folder, so a relative staging path would land there.
    output = io.resolve_output(source, "md2pdf", args.output, suffix="", extension=".pdf").resolve()
    if args.dry_run:
        return _dry_run("md2pdf", output)
    output.parent.mkdir(parents=True, exist_ok=True)
    io.sweep_stale(output.parent)
    problems, check_detail = [], None
    with io.temp_folder() as temp, io.Staging(output, temp) as stage:
        io.disk_budget([source], output, temp, multiplier=3)
        status_path = Path(temp) / "render-status.json"
        env = _node_env(PDFTOOLS_EDGE=edge, PDFTOOLS_HTML=str(Path(temp) / "render.html"),
                        PDFTOOLS_CSS=str(ROOT / "md2pdf.css"), PDFTOOLS_STATUS=str(status_path))
        try:
            result = tools.run_tool([node, str(script), "--render", str(source), str(stage.path)],
                                    timeout=120, cwd=temp, env=env, temp=temp,
                                    secret_env=getattr(context, "secret_env", ()))
        except clierr.Fail as exc:
            raise _fail_tool(exc)
        try:
            status = json.loads(status_path.read_text(encoding="utf-8"))
            if not isinstance(status, dict) or not {"warnings", "errors", "scriptExecuted"}.issubset(status):
                raise ValueError("incomplete render status")
        except (OSError, ValueError, TypeError) as exc:
            raise PdfError("TOOL_FAILED", "Markdown renderer did not write a valid status file",
                           "retry with a working Node renderer", detail=str(exc))
        if status.get("scriptExecuted"):
            raise PdfError("CHECK_FAILED", "Markdown page script executed despite its content security policy",
                           "do not use this conversion until the browser security policy is restored")
        warnings = list(status.get("warnings", []))
        if status.get("errors") or {"KATEX_ERROR", "LOCAL_IMAGE_FAILED"}.intersection(warnings):
            problems.append(("RENDER_INCOMPLETE", "a formula or local image failed to render"))
        try:
            with pymupdf.open(stage.path) as doc:
                pages = doc.page_count
            source_check = check_markdown_source(stage.path, source, context, temp=temp)
        except PdfError as exc:
            if exc.code != "CHECK_FAILED":
                raise
            problems.append(("SOURCE_CHECK_FAILED", exc.what))
            try:
                detail_data = json.loads(exc.detail or "{}")
            except (TypeError, ValueError):
                detail_data = {}
            source_check = {"data": {"warnings": detail_data.get("warnings", []) if isinstance(detail_data, dict) else []}}
            check_detail = exc.detail
        except Exception as exc:
            raise PdfError("TOOL_FAILED", f"could not validate rendered Markdown: {exc}",
                           "check the source and renderer", detail=str(exc))
        source_warnings = source_check.get("data", {}).get("warnings", [])
        warnings.extend(w.get("code", "SOURCE_WARNING") for w in source_warnings)
        warnings.extend(code for code, _ in problems if code not in warnings)
        log = _unchecked("md2pdf", stage, output, context, problems, check_detail) if problems else None
        final = stage.publish()
    if final != stage.target:
        warnings.append("NAME_TAKEN")
    for code, message in _markdown_diagnostics(warnings, source_warnings, problems):
        clierr.warn("pdf md2pdf", code, message, "review the converted document")
    warnings = list(dict.fromkeys(warnings))
    record = _result("md2pdf", final, pages, warnings, verdict="unchecked" if problems else ("warn" if warnings else "ok"))
    if log:
        record["log"] = log
    return record


_TEX = (r"\frac", r"\partial", "^{", "_{")
_BOTTOM_MARGIN = 20 / 25.4 * 72  # md2pdf.css @page margin


def _norm_word(word):
    return re.sub(r"[^\w'-]", "", word, flags=re.UNICODE).casefold()


def _nospace(text):
    return re.sub(r"\s+", "", text)


def check_markdown_source(pdf_path, markdown_path, context, temp=None):
    pdf_path, markdown_path = _path(pdf_path), _clean_markdown_path(markdown_path)
    node = shutil.which("node")
    if not node:
        raise PdfError("NOT_FOUND", "Node.js is not installed", "install Node.js")
    script = ROOT / "md2pdf.js"
    try:
        with (contextlib.nullcontext(temp) if temp else io.temp_folder()) as work:
            entities = Path(work) / "entities.json"
            entities.write_text(json.dumps({k[:-1]: v for k, v in html_entities.html5.items() if k.endswith(";")}),
                                encoding="utf-8")
            result = tools.run_tool([node, str(script), "--tokens", str(markdown_path), str(entities)], timeout=30,
                                    env=_node_env(), cwd=work, temp=work,
                                    secret_env=getattr(context, "secret_env", {}))
        parsed = json.loads(result.stdout)
        blocks, source_images = parsed["blocks"], parsed.get("images", 0)
        with open_input(pdf_path, context.password or context.owner_password) as doc:
            return _compare_source(doc, blocks, source_images, pdf_path, markdown_path)
    except PdfError:
        raise
    except clierr.Fail as exc:
        raise _fail_tool(exc)
    except Exception as exc:
        raise PdfError("TOOL_FAILED", f"Markdown source check failed: {exc}",
                       "check the Markdown and Node installation", detail=str(exc))


def _compare_source(doc, blocks, source_images, pdf_path, markdown_path):
    pdf_words = [(page_no, _norm_word(word[4])) for page_no, page in enumerate(doc, 1)
                 for word in page.get_text("words")]
    pdf_words = [(p, w) for p, w in pdf_words if w]
    pdf_text = _nospace("".join(page.get_text("text") for page in doc))
    # Exempt only code occurrences, consumed in source order; an identical formula elsewhere stays checked.
    scrubbed = list(pdf_text)
    code_cursor = 0
    for block in blocks:
        for code in block.get("code", []):
            piece = _nospace(code)
            if not piece:
                continue
            position = pdf_text.find(piece, code_cursor)
            if position >= 0:
                scrubbed[position:position + len(piece)] = " " * len(piece)
                code_cursor = position + len(piece)
    scrubbed = "".join(scrubbed)
    raw_tex = [tex for tex in _TEX if tex in scrubbed]

    source = []  # (block index, word)
    for index, block in enumerate(blocks):
        text = block.get("text", "")
        source.extend((index, w) for w in (_norm_word(x) for x in text.split()) if w)
    matcher = difflib.SequenceMatcher(None, [w for _, w in source], [w for _, w in pdf_words], autojunk=False)
    matched, consumed = set(), set()
    for a, b, size in matcher.get_matching_blocks():
        matched.update(range(a, a + size))
        consumed.update(range(b, b + size))
    # A very long token (URL) may be broken across PDF lines into several words. Only
    # recover it inside the corresponding gap between already aligned source/PDF spans.
    aligned = [(a, b, size) for a, b, size in matcher.get_matching_blocks() if size]
    for position in range(len(source)):
        token = source[position][1]
        if position in matched or len(token) < 12:
            continue
        previous = [(a + size, b + size) for a, b, size in aligned if a + size <= position]
        following = [(a, b) for a, b, size in aligned if a > position]
        source_floor, pdf_floor = max(previous, default=(0, 0), key=lambda pair: pair[0])
        source_ceiling, pdf_ceiling = min(following, default=(len(source), len(pdf_words)), key=lambda pair: pair[0])
        if not source_floor <= position < source_ceiling:
            continue
        for start in range(pdf_floor, pdf_ceiling):
            if start in consumed:
                continue
            combined = ""
            end = start
            while end < pdf_ceiling and end not in consumed and len(combined) < len(token):
                combined += pdf_words[end][1]
                end += 1
            if combined == token:
                matched.add(position)
                consumed.update(range(start, end))
                break

    per_block = {}
    for position, (index, _) in enumerate(source):
        per_block.setdefault(index, []).append(position in matched)
    missing, warnings = [], []
    for index, flags in per_block.items():
        section = blocks[index].get("section", "document")
        absent = flags.count(False)
        if absent and absent >= len(flags) * .5:
            missing.append({"section": section, "missing": absent, "words": len(flags)})
        run = longest = 0
        for flag in flags:
            run = 0 if flag else run + 1
            longest = max(longest, run)
        if longest >= 3:
            warnings.append({"code": "MISSING_WORD_RUN", "section": section})

    headings = {_nospace("".join(_norm_word(w) for w in b["text"].split())) for b in blocks if b.get("type") == "heading"}
    last_page = doc.page_count
    for page_no, page in enumerate(doc, 1):
        height = page.rect.height
        lines = []
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                text = "".join(span.get("text", "") for span in line.get("spans", []))
                if text.strip():
                    lines.append((line["bbox"], text, line.get("spans", [])))
        # Zero-width struts (KaTeX) are not text anyone reads: count only spans with letters or digits.
        if any(span.get("size", 99) < 6 and any(ch.isalnum() for ch in span.get("text", ""))
               for _, _, spans in lines for span in spans):
            warnings.append({"code": "SMALL_TEXT", "page": page_no})
        if lines and page_no < last_page:
            bbox, text, _ = max(lines, key=lambda item: item[0][3])
            folded = _nospace("".join(_norm_word(w) for w in text.split()))
            if folded and bbox[3] > height * .85 and any(
                    folded == heading or (len(folded) >= 12 and heading.startswith(folded))
                    for heading in headings):
                warnings.append({"code": "HEADING_ORPHAN", "page": page_no})
        limit = height - _BOTTOM_MARGIN + 3
        if any(bbox[3] > limit for bbox, _, _ in lines) or any(
                info["bbox"][3] > limit and info["bbox"][3] - info["bbox"][1] <= height - 2 * _BOTTOM_MARGIN
                for info in page.get_image_info()):
            warnings.append({"code": "BLOCK_CUT", "page": page_no})
    pdf_image_count = sum(len(page.get_images(full=True)) for page in doc)
    if source_images != pdf_image_count:
        warnings.append({"code": "IMAGE_COUNT_CHANGED", "source": source_images, "output": pdf_image_count})
    if missing:
        findings = [{"code": "MISSING_SOURCE_TEXT", "message": f"Markdown blocks missing at least half their words: {', '.join(dict.fromkeys(x['section'] for x in missing))}"}]
        findings.extend({"code": "RAW_TEX", "message": f"raw TeX command remains in PDF text: {tex}"} for tex in raw_tex)
        raise PdfError("CHECK_FAILED",
                       f"Markdown blocks missing at least half their words: {', '.join(dict.fromkeys(x['section'] for x in missing))}",
                       "review the PDF against its Markdown source", detail=json.dumps({"findings": findings, "warnings": warnings}))
    if raw_tex:
        findings = [{"code": "RAW_TEX", "message": f"raw TeX command remains in PDF text: {tex}"} for tex in raw_tex]
        raise PdfError("CHECK_FAILED", findings[0]["message"], "fix the Markdown math delimiters",
                       detail=json.dumps({"findings": findings, "warnings": warnings}))
    return {"type": "report", "command": "check", "input": str(pdf_path), "verdict": "warn" if warnings else "ok",
            "data": {"source": str(markdown_path), "blocks": len(blocks), "warnings": warnings}}


# ---------------------------------------------------------------------------------------------
# Text -> PDF

def add_arguments_txt2pdf(parser):
    parser.add_argument("file", type=Path)
    parser.add_argument("-o", "--output")
    parser.add_argument("--mono", action="store_true")


def txt2pdf(args, context):
    source = _path(args.file)
    try:
        text = source.read_text(encoding="utf-8-sig")
    except UnicodeError as exc:
        raise PdfError("BAD_INPUT", "text file is not valid UTF-8", "save the file as UTF-8", detail=str(exc))
    if not text.strip():
        raise PdfError("BAD_INPUT", "text file is empty", "give a file with text")

    def writer(src, staged, temp):
        font = "monospace" if args.mono else "sans-serif"
        # break-word: a run longer than the line (URL, base64) wraps instead of being clipped.
        story = pymupdf.Story(f'<div style="font-family:{font}; white-space:pre-wrap; overflow-wrap:break-word">'
                              f'{html.escape(text, quote=False)}</div>')
        out = pymupdf.DocumentWriter(str(staged))
        try:
            more = True
            while more:
                page = out.begin_page(pymupdf.Rect(0, 0, 595, 842))
                more, _ = story.place(pymupdf.Rect(56, 56, 539, 786))
                story.draw(page)
                out.end_page()
        finally:
            out.close()

    def verify(pdf):
        # Invariant: every non-space character of the input is in the PDF, in order.
        with pymupdf.open(pdf) as doc:
            got = unicodedata.normalize("NFC", _nospace("".join(page.get_text() for page in doc)))
        want = unicodedata.normalize("NFC", _nospace(text))
        if got == want:
            return []
        complex_scripts = ("ARABIC", "BENGALI", "DEVANAGARI", "GUJARATI", "GURMUKHI", "KANNADA",
                           "KHMER", "LAO", "MALAYALAM", "MYANMAR", "ORIYA", "SINHALA", "TAMIL", "TELUGU", "THAI")
        if any(any(script in unicodedata.name(char, "") for script in complex_scripts) for char in text):
            return {"problems": [], "warnings": [("TEXT_MISMATCH", "complex-script text extraction differs after NFC normalization")]}
        return [("TEXT_MISMATCH", f"PDF holds {len(got)} non-space characters, the text file {len(want)}")]
    return _publish_pdf(source, args, context, writer, command_name="txt2pdf", verify=verify)


office2pdf = command("office2pdf", 6, "any", "pdf", "", False, "report")(office2pdf)
toword = command("toword", 6, "pdf", "docx", "", False, "report")(toword)
md2pdf = command("md2pdf", 6, "markdown", "pdf", "", False, "report")(md2pdf)
txt2pdf = command("txt2pdf", 6, "text", "pdf", "", False, "report")(txt2pdf)


def selftest_write(source, output):
    from docx import Document
    doc = Document()
    doc.add_paragraph("pdftools phase 6 selftest")
    doc.save(output)
