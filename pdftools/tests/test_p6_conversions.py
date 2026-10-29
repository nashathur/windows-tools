import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT.parent / "common"))

import pymupdf
from PIL import Image
from pdftools import cli, tools
from pdftools.commands import p2_selftest, p3c_helpers, p6_convert
from tests.fixtures import sans_font
from tests.fixtures.p6 import (PASSWORD, build_encrypted_pdf, build_fixtures, build_markdown_fixtures,
                               build_office_tripwires)

TOOL_USAGE = {}


def tool_info(name):
    if name == "pdf2docx":
        from importlib.metadata import PackageNotFoundError, version
        try:
            return "pdf2docx", version("pdf2docx")
        except PackageNotFoundError as exc:
            return None, f"pdf2docx is not installed: {exc}"
    path = tools.find_tool(name) if name in tools.TOOL_CANDIDATES else shutil.which(name)
    if not path:
        return None, f"{name} executable is absent"
    if Path(path).name.lower() == "msedge.exe":
        command = "(Get-Item -LiteralPath '" + path.replace("'", "''") + "').VersionInfo.ProductVersion"
        try:
            proc = subprocess.run(["powershell.exe", "-NoProfile", "-Command", command], capture_output=True, text=True, timeout=8)
            value = proc.stdout.strip().splitlines()
            return path, value[0] if value else f"exit {proc.returncode}"
        except Exception as exc:
            return None, str(exc)
    version_args = ["-version"] if name == "ghostscript" else ["--version"]
    try:
        proc = subprocess.run([path, *version_args], capture_output=True, text=True, timeout=8)
        value = (proc.stdout or proc.stderr).strip().splitlines()
        return path, value[0] if value else f"exit {proc.returncode}"
    except Exception as exc:
        return None, str(exc)


class Phase6Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        cls.fixtures = build_fixtures(cls.root / "fixtures")

    @classmethod
    def tearDownClass(cls):
        for name, item in sorted(TOOL_USAGE.items()):
            print(f"TOOL {name}: {item['version']}; tests run {len(item['tests'])}", file=sys.stderr)
        cls.temp.cleanup()

    def run_cli(self, argv):
        out, err = io.StringIO(), io.StringIO()
        with mock.patch("clierr.LOG_ROOT", self.root / "logs"), contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = cli.main(["--json", "--stdout", *map(str, argv)])
        rows = [json.loads(line) for line in out.getvalue().splitlines() if line.strip()]
        return code, rows, err.getvalue()

    def require_tool(self, name):
        path, version = tool_info(name)
        if not path:
            reason = version or "executable is absent"
            if os.environ.get("PDFTOOLS_REQUIRE_TOOLS") == "1":
                self.fail(f"required tool {name} is missing: {reason}")
            print(f"SKIP {name}: {reason}", file=sys.stderr)
            self.skipTest(reason)
        usage = TOOL_USAGE.setdefault(name, {"version": version, "tests": set()})
        usage["tests"].add(self.id())
        return path, version

    def test_selftest_missing_converter_packages_are_failures(self):
        failures = p2_selftest._missing_library_failures({
            "puppeteer-core": "absent: missing", "katex": "absent: missing", "marked": "absent: missing",
            "pdf2docx": "absent: missing", "docx": "absent: missing",
        })
        self.assertEqual(len(failures), 5, failures)
        self.assertTrue(any("python-docx" in failure for failure in failures))

    def test_md2pdf_source_warnings_deduplicate_and_keep_page_section(self):
        diagnostics = p6_convert._markdown_diagnostics(
            ["SMALL_TEXT", "SMALL_TEXT"],
            [{"code": "SMALL_TEXT", "page": 2, "section": "Rainfall"},
             {"code": "SMALL_TEXT", "page": 2, "section": "Rainfall"}], [],
        )
        self.assertEqual(diagnostics, [("SMALL_TEXT", "Markdown conversion warning: page 2, section Rainfall")])

    def test_source_check_split_token_recovery_stays_between_aligned_words(self):
        pdf = self.make_pdf("reordered-token.pdf", [[(50, 100, "before abcdefghijklmnop", 12)]])
        with pymupdf.open(pdf) as doc, self.assertRaises(p6_convert.PdfError) as caught:
            p6_convert._compare_source(doc, [{"type": "paragraph", "section": "order",
                                               "text": "abcdefghijklmnop before", "code": []}], 0, pdf, pdf)
        self.assertEqual(caught.exception.code, "CHECK_FAILED")

    def test_source_check_does_not_decode_html_entities_twice_in_code(self):
        pdf = self.make_pdf("literal-code.pdf", [[]])
        with pymupdf.open(pdf) as doc, self.assertRaises(p6_convert.PdfError) as caught:
            p6_convert._compare_source(doc, [{"type": "code", "section": "code",
                                               "text": "&amp;", "code": ["&amp;"]}], 0, pdf, pdf)
        self.assertEqual(caught.exception.code, "CHECK_FAILED")

    def test_img2pdf_images_info_diff_and_ocr(self):
        self.require_tool("tesseract")
        output = self.root / "chart.pdf"
        code, rows, _ = self.run_cli(["img2pdf", self.fixtures["image"], "-o", output])
        self.assertEqual(code, 0, rows)
        with pymupdf.open(output) as doc:
            self.assertEqual(doc.page_count, 1)
            self.assertEqual(len(doc[0].get_images()), 1)
        code, rows, _ = self.run_cli(["imginfo", self.fixtures["image"]])
        self.assertEqual(code, 0)
        self.assertEqual(rows[0]["data"]["size"], "400x250")
        code, rows, _ = self.run_cli(["imgdiff", self.fixtures["image"], self.fixtures["changed_image"]])
        self.assertEqual(code, 0)
        self.assertIn("data", rows[0], rows)
        self.assertTrue(rows[0]["data"]["findings"])
        code, rows, _ = self.run_cli(["imgtext", self.fixtures["image"], "--lang", "eng"])
        self.assertEqual(code, 0, rows)
        self.assertIn("weather", rows[0]["data"]["text"].lower())  # the fixture reads "Weather 2026"
        self.assertGreaterEqual(rows[0]["data"]["words"], 2)

    def test_toword_and_md2pdf_sweep_stale_staging_before_writing(self):
        self.require_tool("pdf2docx")
        self.require_markdown_tools()
        folder = self.root / "sweep"
        folder.mkdir()
        for command_name, source, name in (("toword", self.fixtures["pdf"], "w.docx"), ("md2pdf", self.fixtures["markdown"], "m.pdf")):
            stale = folder / f".{command_name}.2147483000-deadbeef.part{Path(name).suffix}"
            stale.write_bytes(b"leftover of a crashed run")
            old = time.time() - 7200
            os.utime(stale, (old, old))
            code, rows, err = self.run_cli([command_name, source, "-o", folder / name])
            self.assertEqual(code, 0, (rows, err))
            self.assertFalse(stale.exists(), command_name)

    def test_img2pdf_eps_uses_ghostscript_route(self):
        self.require_tool("ghostscript")
        output = self.root / "eps.pdf"
        code, rows, _ = self.run_cli(["img2pdf", self.fixtures["eps"], "-o", output])
        self.assertEqual(code, 0, rows)
        with pymupdf.open(output) as doc:
            self.assertEqual(doc.page_count, 1)

    def test_images_extracts_xobjects_and_png_can_be_opened(self):
        output = self.root / "extracted"
        code, rows, _ = self.run_cli(["images", self.fixtures["pdf"], "-o", str(output) + "\\"])
        self.assertEqual(code, 0, rows)
        self.assertTrue((Path(rows[0]["path"]) / "images.json").is_file())

    def test_toimg_additional_writers_and_validation(self):
        cases = (("webp", "WEBP"), ("tiff", "TIFF"), ("bmp", "BMP"), ("gif", "GIF"), ("avif", "AVIF"), ("heic", "HEIF"))
        for fmt, expected in cases:
            with self.subTest(fmt=fmt):
                target = self.root / f"page.{fmt}"
                code, rows, _ = self.run_cli(["toimg", self.fixtures["pdf"], "--format", fmt, "--pages", "1", "-o", target])
                self.assertEqual(code, 0, rows)
                with Image.open(target) as image:
                    self.assertEqual(image.format, expected)
        code, rows, _ = self.run_cli(["toimg", self.fixtures["pdf"], "--format", "avif", "--lossless", "--pages", "1"])
        self.assertEqual(code, 2)
        self.assertEqual(rows[0]["code"], "BAD_ARG")

    # ---- helpers -------------------------------------------------------------------------------

    def pdf_text(self, path):
        with pymupdf.open(path) as doc:
            return " ".join(page.get_text() for page in doc)

    def write_md(self, name, text):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def require_markdown_tools(self):
        self.require_tool("node")
        self.require_tool("browser")

    def make_pdf(self, name, pages):
        """pages: list of lists of (x, y, text, fontsize) -> hand-made PDF (no toolkit code)."""
        path = self.root / name
        doc = pymupdf.open()
        for items in pages:
            page = doc.new_page()
            for x, y, text, size in items:
                page.insert_text((x, y), text, fontsize=size)
        doc.save(path)
        doc.close()
        return path

    def source_check(self, pdf, md):
        code, rows, err = self.run_cli(["check", pdf, "--source", md])
        return code, rows[0], err

    @contextlib.contextmanager
    def junction(self, link, target):
        if os.name != "nt":  # POSIX has no junctions; a directory symlink is the same escape
            os.symlink(target, link, target_is_directory=True)
        else:
            proc = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(target)], capture_output=True, text=True)
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        try:
            yield
        finally:
            (os.unlink if os.name != "nt" else os.rmdir)(link)  # removes the link only, never what it points to

    # ---- txt2pdf -------------------------------------------------------------------------------

    def test_txt2pdf_clean_text_is_ok_and_escaped(self):
        output = self.root / "text.pdf"
        code, rows, _ = self.run_cli(["txt2pdf", self.fixtures["text"], "-o", output])
        self.assertEqual(code, 0, rows)
        self.assertEqual(rows[0]["verdict"], "ok")
        self.assertIn("A & B < C > D", self.pdf_text(output))

    def test_txt2pdf_long_unbroken_run_is_wrapped_not_clipped(self):
        source = self.write_md("long.txt", "start " + "x" * 20000 + " end\n" + "y" * 5000 + "\n")
        output = self.root / "long.pdf"
        code, rows, _ = self.run_cli(["txt2pdf", source, "-o", output])
        self.assertEqual(code, 0, rows)
        self.assertEqual(rows[0]["verdict"], "ok")
        self.assertEqual("".join(self.pdf_text(output).split()), "".join(source.read_text(encoding="utf-8").split()))

    def test_txt2pdf_devanagari_conjuncts_compare_after_nfc(self):
        source = self.write_md("devanagari.txt", "नमस्ते दुनिया\n")
        code, rows, _ = self.run_cli(["txt2pdf", source, "-o", self.root / "devanagari.pdf"])
        self.assertEqual((code, rows[0]["verdict"], rows[0]["warnings"]),
                         (0, "warn", ["TEXT_MISMATCH"]), rows)

    def test_txt2pdf_clipped_output_is_unchecked(self):
        """The invariant (every non-space character present) catches a clipping Story."""
        source = self.write_md("clip.txt", "z" * 3000 + "\n")
        real_story = p6_convert.pymupdf.Story

        def clipping(html_text, *a, **kw):
            return real_story(html_text.replace("overflow-wrap:break-word", ""), *a, **kw)
        with mock.patch.object(p6_convert.pymupdf, "Story", clipping):
            code, rows, err = self.run_cli(["txt2pdf", source, "-o", self.root / "clip.pdf"])
        self.assertEqual(code, 3, rows)
        self.assertEqual(rows[0]["verdict"], "unchecked")
        self.assertTrue(rows[0]["path"].endswith("clip_UNCHECKED.pdf"), rows)
        self.assertIn("WARN pdf txt2pdf TEXT_MISMATCH", err)

    def test_txt2pdf_empty_file_is_bad_input(self):
        code, rows, _ = self.run_cli(["txt2pdf", self.write_md("empty.txt", "  \n"), "-o", self.root / "empty.pdf"])
        self.assertEqual((code, rows[0]["code"]), (1, "BAD_INPUT"))

    def test_txt2pdf_dry_run_creates_nothing(self):
        output = self.root / "dry.pdf"
        code, rows, _ = self.run_cli(["--dry-run", "txt2pdf", self.fixtures["text"], "-o", output])
        self.assertEqual((code, rows[0]["verdict"], rows[0]["provisional"]), (0, "report", True))
        self.assertFalse(output.exists())

    # ---- office2pdf ----------------------------------------------------------------------------

    def test_office2pdf_converts_and_reports_without_claiming_a_check(self):
        self.require_tool("soffice")
        output = self.root / "office.pdf"
        code, rows, _ = self.run_cli(["office2pdf", self.fixtures["docx"], "-o", output])
        self.assertEqual(code, 0, rows)
        self.assertEqual(rows[0]["verdict"], "report")
        self.assertIn("Office conversion", self.pdf_text(output))

    def test_office2pdf_profile_keys_exist_in_the_installed_libreoffice_schema(self):
        import xml.etree.ElementTree as ET
        soffice, _ = self.require_tool("soffice")
        registry = Path(soffice).resolve().parent.parent / "share" / "registry" / "main.xcd"
        oor = "{http://openoffice.org/2001/registry}"
        schema = ET.parse(registry).getroot()

        def schema_default(path, prop):
            package, _, name = path.strip("/").split("/")[0].rpartition(".")
            for component in schema.iter(oor + "component-schema"):
                if component.get(oor + "name") != name or component.get(oor + "package") != package:
                    continue
                nodes = [component.find("component")]
                for part in path.strip("/").split("/")[1:]:
                    nodes = [child for node in nodes for child in node if child.get(oor + "name") == part]
                for node in nodes:
                    for child in node:
                        if child.tag == "prop" and child.get(oor + "name") == prop:
                            return True
            return False
        profile = self.root / "schema-profile"
        p6_convert._office_security(profile)
        items = ET.parse(profile / "user" / "registrymodifications.xcu").getroot()
        written = {}
        for item in items:
            for prop in item:
                self.assertTrue(schema_default(item.get(oor + "path"), prop.get(oor + "name")),
                                f"{item.get(oor + 'path')} {prop.get(oor + 'name')} is not in LibreOffice's schema")
                written[(item.get(oor + "path"), prop.get(oor + "name"))] = prop.find("value").text
        self.assertEqual(written[("/org.openoffice.Office.Common/Security/Scripting", "MacroSecurityLevel")], "3")
        self.assertEqual(written[("/org.openoffice.Office.Common/Security/Scripting", "DisableMacrosExecution")], "true")
        self.assertEqual(written[("/org.openoffice.Office.Common/Security/Scripting", "BlockUntrustedRefererLinks")], "true")
        self.assertEqual(written[("/org.openoffice.Office.Writer/Content/Update", "Link")], "0")
        self.assertEqual(written[("/org.openoffice.Office.Calc/Content/Update", "Link")], "1")
        self.assertNotIn(("/org.openoffice.Office.Common/Content/Update", "Link"), written)

    def macro_fires(self, soffice, profile, document, marker, wait):
        """Open `document` (visible-style load, which is when document events fire) and watch for the marker."""
        import psutil
        marker.unlink(missing_ok=True)
        proc = subprocess.Popen([soffice, "--headless", "--norestore", "--nolockcheck",
                                 f"-env:UserInstallation={Path(profile).as_uri()}", str(document)],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            deadline = time.monotonic() + wait
            while time.monotonic() < deadline and not marker.exists():
                time.sleep(0.25)
            return marker.exists()
        finally:
            try:
                parent = psutil.Process(proc.pid)
                for child in parent.children(recursive=True):
                    child.kill()
                parent.kill()
            except psutil.Error:
                pass

    def test_office2pdf_macro_and_link_tripwires(self):
        soffice, _ = self.require_tool("soffice")
        fx = build_office_tripwires(self.root / "tripwire")
        permissive = self.root / "permissive-profile"
        (permissive / "user").mkdir(parents=True)
        (permissive / "user" / "registrymodifications.xcu").write_text(
            '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry" '
            'xmlns:xs="http://www.w3.org/2001/XMLSchema"><item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
            '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>0</value></prop>'
            '<prop oor:name="CheckDocumentEvents" oor:op="fuse"><value>false</value></prop></item></oor:items>', encoding="utf-8")
        self.assertTrue(self.macro_fires(soffice, permissive, fx["macro_odt"], fx["marker"], 60),
                        "control: the fixture's open macro must run under a permissive profile")
        ours = self.root / "ours-profile"
        p6_convert._office_security(ours)
        self.assertFalse(self.macro_fires(soffice, ours, fx["macro_odt"], fx["marker"], 15),
                         "the open macro ran under the office2pdf profile")
        # No positive control for linked content: headless `--convert-to` never refreshes document links,
        # even under an update-always profile (verified on this LibreOffice), so only the stale-cache
        # assertions below can be made; the macro half above keeps its positive control.
        fx["marker"].unlink(missing_ok=True)
        code, rows, _ = self.run_cli(["office2pdf", fx["macro_odt"], "-o", self.root / "macro.pdf"])
        self.assertEqual(code, 0, rows)
        self.assertFalse(fx["marker"].exists(), "macro ran during conversion")
        code, rows, _ = self.run_cli(["office2pdf", fx["link_odt"], "-o", self.root / "link.pdf"])
        self.assertEqual(code, 0, rows)
        text = self.pdf_text(self.root / "link.pdf")
        self.assertIn("STALE-CACHED-TEXT", text)
        self.assertNotIn("FRESH-LINKED-TEXT", text)

    # ---- toword --------------------------------------------------------------------------------

    def test_toword_clean_pdf_is_ok(self):
        self.require_tool("pdf2docx")
        output = self.root / "word.docx"
        code, rows, _ = self.run_cli(["toword", self.fixtures["pdf"], "-o", output])
        self.assertEqual(code, 0, rows)
        self.assertEqual((rows[0]["verdict"], rows[0]["warnings"]), ("ok", []))
        from docx import Document
        self.assertIn("Phase six", " ".join(p.text for p in Document(output).paragraphs))

    def test_toword_password_never_in_argv_or_child_env_and_export_is_warned(self):
        self.require_tool("pdf2docx")
        locked = build_encrypted_pdf(self.root / "enc")
        pw_file = self.root / "pw.txt"
        pw_file.write_text(PASSWORD + "\n", encoding="utf-8")
        seen = []
        real_popen = subprocess.Popen

        def spy(argv, *a, **kw):
            seen.append((list(map(str, argv)), dict(kw.get("env") if kw.get("env") is not None else os.environ)))
            return real_popen(argv, *a, **kw)
        forms = {PASSWORD, json.dumps(PASSWORD), json.dumps(PASSWORD)[1:-1], repr(PASSWORD)}
        for mode, flags in (("file", ["--password-file", pw_file]), ("env", ["--password-env", "P6_TEST_PASSWORD"])):
            output = self.root / f"locked-{mode}.docx"
            seen.clear()
            with mock.patch.dict(os.environ, {"P6_TEST_PASSWORD": PASSWORD} if mode == "env" else {}), mock.patch.object(subprocess, "Popen", spy):
                code, rows, err = self.run_cli(["toword", locked, "-o", output, *flags])
            self.assertEqual(code, 0, (mode, rows, err))
            self.assertTrue(seen, "no child process was started")
            for argv, env in seen:
                for form in forms:
                    self.assertFalse(any(form in arg for arg in argv), (mode, argv))
                    self.assertFalse(any(form in str(value) for value in env.values()), mode)
                self.assertNotIn("P6_TEST_PASSWORD", env)
            self.assertIn("DECRYPTED_EXPORT", rows[0]["warnings"])
            self.assertEqual(rows[0]["verdict"], "warn")
            self.assertIn("WARN pdf toword DECRYPTED_EXPORT", err)
            from docx import Document
            self.assertIn("Encrypted conversion", " ".join(p.text for p in Document(output).paragraphs))
        wrong = self.root / "wrong.txt"
        wrong.write_text("not the password\n", encoding="utf-8")
        code, rows, err = self.run_cli(["toword", locked, "-o", self.root / "wrong.docx", "--password-file", wrong])
        self.assertEqual((code, rows[0]["code"]), (1, "WRONG_PASSWORD"))
        leaked = [str(path) for path in (self.root / "logs").rglob("*") if path.is_file()
                  and any(form in path.read_text(encoding="utf-8", errors="replace") for form in forms)]
        self.assertEqual(leaked, [])

    def test_toword_secret_environment_is_filtered_without_pdf2docx(self):
        output = self.root / "mock-word.docx"
        seen = {}
        real_run_tool = p6_convert.tools.run_tool
        def warned(*args, **kwargs):
            seen.setdefault("warnings", []).append(args[1])
        def fake_run_tool(argv, **kwargs):
            seen.update(kwargs)
            real_run_tool([sys.executable, "-c", "import os,sys; sys.exit('P6_SECRET' in os.environ)"], **kwargs)
            target = json.loads(Path(argv[3]).read_text(encoding="utf-8"))[1]
            from docx import Document
            doc = Document(); doc.add_paragraph("few"); doc.save(target)
        context = mock.Mock(password=None, owner_password=None, secret_env=["P6_SECRET"],
                            secret_store=mock.Mock(values=[]))
        with mock.patch.dict(os.environ, {"P6_SECRET": "keep-out-of-child"}), \
                mock.patch.object(p6_convert, "_signed", return_value=False), \
                mock.patch.object(p6_convert.tools, "run_tool", side_effect=fake_run_tool), \
                mock.patch.object(p6_convert.clierr, "warn", side_effect=warned):
            p6_convert.toword(mock.Mock(file=self.fixtures["pdf"], output=output, dry_run=False), context)
        self.assertEqual(seen["secret_env"], ["P6_SECRET"])
        self.assertEqual(seen["warnings"].count("LOW_TEXT"), 1)

    def test_source_check_child_receives_secret_environment_exclusions_without_optional_tools(self):
        context = mock.Mock(password=None, owner_password=None, secret_env={"P6_TEST_PASSWORD": None})
        result = mock.Mock(stdout=json.dumps({"blocks": [], "images": 0}))
        with mock.patch("shutil.which", return_value="node.exe"), \
                mock.patch.object(p6_convert.tools, "run_tool", return_value=result) as run:
            p6_convert.check_markdown_source(self.fixtures["pdf"], self.fixtures["markdown"], context)
        self.assertEqual(run.call_args.kwargs["secret_env"], {"P6_TEST_PASSWORD": None})

    def test_toword_unencrypted_pdf_has_no_export_warning(self):
        self.require_tool("pdf2docx")
        code, rows, _ = self.run_cli(["toword", self.fixtures["pdf"], "-o", self.root / "plain.docx"])
        self.assertNotIn("DECRYPTED_EXPORT", rows[0]["warnings"])

    # ---- md2pdf --------------------------------------------------------------------------------

    def test_md2pdf_clean_fixture_is_ok_and_check_source_agrees(self):
        self.require_markdown_tools()
        output = self.root / "markdown.pdf"
        code, rows, err = self.run_cli(["md2pdf", self.fixtures["markdown"], "-o", output])
        self.assertEqual((code, rows[0]["verdict"], rows[0]["warnings"]), (0, "ok", []), (rows, err))
        code, record, _ = self.source_check(output, self.fixtures["markdown"])
        self.assertEqual((code, record["verdict"]), (0, "ok"), record)

    def test_md2pdf_relative_output_path_works(self):
        self.require_markdown_tools()
        previous = os.getcwd()
        os.chdir(self.root)
        try:
            code, rows, err = self.run_cli(["md2pdf", self.fixtures["markdown"], "-o", "rel-out.pdf"])
        finally:
            os.chdir(previous)
        self.assertEqual(code, 0, (rows, err))
        self.assertTrue((self.root / "rel-out.pdf").is_file())

    def test_md2pdf_code_renders_as_code_never_as_markup_or_math(self):
        self.require_markdown_tools()
        rich = build_markdown_fixtures(self.root / "md-rich")["rich_md"]
        output = self.root / "rich.pdf"
        code, rows, err = self.run_cli(["md2pdf", rich, "-o", output])
        self.assertEqual((code, rows[0]["verdict"], rows[0]["warnings"]), (0, "ok", []), (rows, err))
        text = self.pdf_text(output)
        flat = "".join(text.split())
        for literal in ("\\frac{1}{2}", '<imgsrc="https://127.0.0.1:9/x.png">', "$$notmath$$", "<b>literal</b>",
                        "$dollars$", "<i>tags</i>", "Costs$5and$6"):
            self.assertIn("".join(literal.split()), flat)
        with pymupdf.open(output) as doc:
            self.assertEqual(sum(len(page.get_images()) for page in doc), 1)  # only pic.png; no <img> from code
        code, record, _ = self.source_check(output, rich)
        self.assertEqual((code, record["verdict"], record["data"]["warnings"]), (0, "ok", []), record)

    def test_md2pdf_equations_use_the_katex_stylesheet_and_print_once(self):
        self.require_markdown_tools()
        md = self.write_md("math.md", "# Math\n\nInline $\\frac{1}{2}$ and display:\n\n$$\\frac{a}{b}=c$$\n")
        output = self.root / "math.pdf"
        code, rows, err = self.run_cli(["md2pdf", md, "-o", output])
        self.assertEqual((code, rows[0]["verdict"]), (0, "ok"), (rows, err))
        with pymupdf.open(output) as doc:
            fonts = " ".join(font[3] for page in doc for font in page.get_fonts())
        self.assertIn("KaTeX", fonts)
        self.assertNotIn("1 2 21", " ".join(self.pdf_text(output).split()))
        self.assertNotIn("\\frac", self.pdf_text(output))

    def test_md2pdf_byte_order_mark_does_not_hide_the_heading(self):
        self.require_markdown_tools()
        bom = build_markdown_fixtures(self.root / "md-bom")["bom_md"]
        output = self.root / "bom.pdf"
        code, rows, err = self.run_cli(["md2pdf", bom, "-o", output])
        self.assertEqual((code, rows[0]["verdict"]), (0, "ok"), (rows, err))
        with pymupdf.open(output) as doc:
            sizes = {span["text"].strip(): span["size"] for block in doc[0].get_text("dict")["blocks"]
                     for line in block.get("lines", []) for span in line["spans"]}
        self.assertGreater(sizes["Heading after BOM"], 15)
        self.assertNotIn("#", self.pdf_text(output))

    def test_md2pdf_wide_equations_fit_the_page_and_give_back_the_space(self):
        self.require_markdown_tools()
        md = build_markdown_fixtures(self.root / "md-eq")["equations_md"]
        output = self.root / "equations.pdf"
        code, rows, err = self.run_cli(["md2pdf", md, "-o", output])
        self.assertEqual(code, 0, (rows, err))
        self.assertIn("EQUATION_SHRUNK", rows[0]["warnings"])
        with pymupdf.open(output) as doc:
            words = doc[0].get_text("words")
            page_width = doc[0].rect.width
        right_margin = page_width - 20 / 25.4 * 72
        self.assertLessEqual(max(w[2] for w in words), right_margin + 1, "equation overflows the printable width")
        after = next(w for w in words if w[4] == "AFTERWORD")
        above = max(w[3] for w in words if w[3] < after[1] - 2)
        self.assertLess(after[1] - above, 25, "shrunk equation left a blank gap")

    def test_md2pdf_security_barrier_blocks_outside_files_and_remote_urls(self):
        self.require_markdown_tools()
        work = self.root / "barrier"
        site, outside = work / "site", work / "outside"
        site.mkdir(parents=True)
        outside.mkdir()
        Image.new("RGB", (30, 30), (0, 0, 255)).save(site / "ok.png")
        Image.new("RGB", (30, 30), (255, 0, 0)).save(outside / "secret.png")
        md = site / "doc.md"
        md.write_text(
            "# Barrier\n\nA paragraph that must survive the conversion untouched by blocked resources.\n\n"
            "![ok](ok.png)\n\n![jn](jn/secret.png)\n\n![up](../outside/secret.png)\n\n"
            f"![abs](<{(outside / 'secret.png').as_uri()}>)\n\n![unc](file://127.0.0.1/share/x.png)\n\n"
            '<img src="https://127.0.0.1:9/remote.png">\n'
            "<script>document.body.dataset.p6Attack='ran'</script>\n", encoding="utf-8")
        with self.junction(site / "jn", outside):
            code, rows, err = self.run_cli(["md2pdf", md, "-o", work / "barrier.pdf"])
        self.assertEqual(code, 3, (rows, err))  # blocked images are failed local images: unchecked
        record = rows[0]
        self.assertEqual(record["verdict"], "unchecked")
        self.assertIn("BLOCKED_LOCAL", record["warnings"])
        self.assertIn("REMOTE_RESOURCE", record["warnings"])
        self.assertNotEqual(record.get("code"), "CHECK_FAILED")  # the page script never ran
        with pymupdf.open(record["path"]) as doc:
            self.assertEqual(sum(len(page.get_images()) for page in doc), 1, "only ok.png may be embedded")
        self.assertTrue((outside / "secret.png").is_file())  # the junction target survived cleanup

    def test_md2pdf_meta_refresh_cannot_run_sibling_html_script(self):
        self.require_markdown_tools()
        md = self.write_md("nav.md", '<meta http-equiv="refresh" content="0;url=sibling.html">\n')
        (md.parent / "sibling.html").write_text('<!doctype html><title>safe</title><script>document.title="p6-nav-marker"</script>', encoding="utf-8")
        code, rows, err = self.run_cli(["md2pdf", md, "-o", self.root / "nav.pdf"])
        self.assertEqual(code, 0, (rows, err))
        self.assertNotIn("script executed", err.lower())
        self.assertNotIn("p6-nav-marker", rows[0].get("warnings", []))
        self.assertIn("BLOCKED_LOCAL", rows[0]["warnings"])

    def test_md2pdf_remote_reparse_target_is_rejected_before_realpath(self):
        self.require_tool("node")
        link = self.root / "remote-link"  # no real reparse point: creating one needs a privilege
        script = ("const fs=require('fs'); const mod=require(process.argv[1]); "
                  "const real=fs.lstatSync; "
                  "fs.lstatSync=(p,...a)=>/remote-link$/.test(String(p))?{isSymbolicLink:()=>true}:real(p,...a); "
                  "fs.readlinkSync=()=>process.argv[3]; "
                  "fs.realpathSync.native=()=>{throw new Error('realpath was reached')}; "
                  "try { mod.rejectRemoteReparseTargets(process.argv[2]); process.exitCode=9; } "
                  "catch(e) { if (!/remote reparse target/.test(e.message)) process.exitCode=8; }")
        result = subprocess.run([shutil.which("node"), "-e", script, str(ROOT / "md2pdf.js"), str(link),
                                 r"\\localhost\c$\does-not-exist"], capture_output=True, text=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_md2pdf_remote_reparse_chain_is_rejected(self):
        self.require_tool("node")
        first = self.root / "link-a"
        script = ("const fs=require('fs'); const mod=require(process.argv[1]); const real=fs.lstatSync; "
                  "fs.lstatSync=(p,...a)=>/link-[ab]$/.test(String(p))?{isSymbolicLink:()=>true}:real(p,...a); "
                  "fs.readlinkSync=p=>String(p).endsWith('link-a')?'link-b':process.argv[3]; "
                  "try { mod.rejectRemoteReparseTargets(process.argv[2]); process.exitCode=9; } "
                  "catch(e) { if (!/remote reparse target/.test(e.message)) process.exitCode=8; }")
        result = subprocess.run([shutil.which("node"), "-e", script, str(ROOT / "md2pdf.js"), str(first),
                                 r"\\server\share"],
                                capture_output=True, text=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_md2pdf_tokens_keep_a_space_at_hard_break(self):
        node, _ = self.require_tool("node")
        md = self.write_md("hard-break.md", "one  \nthree\n")
        result = subprocess.run([node, str(ROOT / "md2pdf.js"), "--tokens", str(md)],
                                capture_output=True, text=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("one three", json.loads(result.stdout)["blocks"][0]["text"])

    def test_md2pdf_remote_only_is_a_warning_and_code_text_is_not_a_remote_resource(self):
        self.require_markdown_tools()
        remote = self.write_md("remote.md", "# Remote\n\nThe climate conversion keeps every word of this paragraph.\n\n"
                                            '<img src="https://127.0.0.1:9/remote.png">\n')
        code, rows, err = self.run_cli(["md2pdf", remote, "-o", self.root / "remote.pdf"])
        self.assertEqual((code, rows[0]["verdict"]), (0, "warn"), (rows, err))
        self.assertEqual(rows[0]["warnings"], ["REMOTE_RESOURCE"])
        quoted = self.write_md("quoted.md", "# Quoted\n\nWrite `src=\"https://example.org/x.png\"` to embed a picture in your notes.\n")
        code, rows, err = self.run_cli(["md2pdf", quoted, "-o", self.root / "quoted.pdf"])
        self.assertEqual((code, rows[0]["verdict"], rows[0]["warnings"]), (0, "ok", []), (rows, err))

    def test_md2pdf_katex_error_is_unchecked(self):
        self.require_markdown_tools()
        md = self.write_md("bad-math.md", "# Bad\n\nParagraph with enough words to be compared properly, then $\\undefinedcommand{x}$ here.\n")
        code, rows, err = self.run_cli(["md2pdf", md, "-o", self.root / "bad-math.pdf"])
        self.assertEqual((code, rows[0]["verdict"]), (3, "unchecked"), (rows, err))
        self.assertTrue(rows[0]["path"].endswith("_UNCHECKED.pdf"))

    def test_md2pdf_source_check_failure_keeps_its_reason(self):
        self.require_markdown_tools()
        reason = "Markdown blocks missing at least half their words: Probe"
        with mock.patch.object(p6_convert, "check_markdown_source",
                               side_effect=p6_convert.PdfError("CHECK_FAILED", reason, "review", detail="[probe-detail]")):
            code, rows, err = self.run_cli(["md2pdf", self.fixtures["markdown"], "-o", self.root / "reason.pdf"])
        self.assertEqual((code, rows[0]["verdict"]), (3, "unchecked"), (rows, err))
        self.assertIn("SOURCE_CHECK_FAILED", rows[0]["warnings"])
        self.assertIn(f"WARN pdf md2pdf SOURCE_CHECK_FAILED: {reason}", err)
        log = Path(rows[0]["log"]).read_text(encoding="utf-8")
        self.assertIn(reason, log)
        self.assertIn("probe-detail", log)

    def test_md2pdf_engine_runs_with_the_browser_sandbox_on(self):
        self.assertNotIn("--no-sandbox", (ROOT / "md2pdf.js").read_text(encoding="utf-8"))

    def test_md2pdf_missing_status_file_fails_closed(self):
        with mock.patch("shutil.which", return_value="node.exe"), mock.patch.object(p6_convert.tools, "find_tool", return_value="edge.exe"), \
                mock.patch.object(p6_convert.tools, "run_tool", return_value=subprocess.CompletedProcess([], 0, "", "")):
            code, rows, _ = self.run_cli(["md2pdf", self.fixtures["markdown"], "-o", self.root / "missing-status.pdf"])
        self.assertEqual((code, rows[0]["code"]), (1, "TOOL_FAILED"), rows)

    # ---- check --source ------------------------------------------------------------------------

    def test_check_source_detects_a_missing_block(self):
        self.require_tool("node")
        pdf = self.root / "check-text.pdf"
        self.run_cli(["txt2pdf", self.fixtures["text"], "-o", pdf])
        code, rows, _ = self.run_cli(["check", pdf, "--source", self.fixtures["markdown"]])
        self.assertEqual((code, rows[0]["verdict"]), (0, "warn"), rows)
        self.assertTrue(rows[0]["data"]["findings"])

    def test_check_source_finding_is_warn_with_exit_zero(self):
        self.require_tool("node")
        pdf = self.make_pdf("source-warning.pdf", [[(50, 100, "only one source word", 11)]])
        md = self.write_md("source-warning.md", "A missing block has enough distinct words to produce a finding.")
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual((code, record["verdict"]), (0, "warn"), record)
        self.assertTrue(record["data"]["findings"], record)

    def test_check_source_compares_code_blocks_instead_of_omitting_them(self):
        self.require_tool("node")
        md = self.write_md("code-only.md", "```text\nclimate_code_marker\n```\n")
        pdf = self.make_pdf("blank-code.pdf", [[]])
        code, rows, _ = self.run_cli(["check", pdf, "--source", md])
        self.assertEqual((code, rows[0]["verdict"]), (0, "warn"), rows)
        self.assertTrue(rows[0]["data"]["findings"], rows)

    def test_check_source_html_text_is_not_an_empty_baseline(self):
        self.require_tool("node")
        md = self.write_md("html.md", "<p>Important HTML paragraph survives.</p>\n")
        pdf = self.make_pdf("blank-html.pdf", [[]])
        code, rows, _ = self.run_cli(["check", pdf, "--source", md])
        self.assertEqual((code, rows[0]["verdict"]), (0, "warn"), rows)

    def test_check_source_decodes_greek_named_entities(self):
        self.require_tool("node")
        md = self.write_md("greek.md", "&alpha; &beta; &gamma;\n")
        pdf = self.root / "greek.pdf"
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_font(fontname="TestFont", fontfile=sans_font())  # base-14 fonts lack Greek
        page.insert_text((50, 100), "α β γ", fontname="TestFont", fontsize=12)
        doc.save(pdf)
        doc.close()
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual(code, 0, record)
        self.assertEqual(record["verdict"], "ok", record)

    def test_check_source_does_not_excuse_a_formula_matching_a_code_occurrence(self):
        self.require_tool("node")
        md = self.write_md("duplicate-tex.md", "`\\frac{1}{2}`\n\nThe rendered equation follows: $\\frac{1}{2}$.\n")
        pdf = self.make_pdf("duplicate-tex.pdf", [[(50, 100, "\\frac{1}{2} \\frac{1}{2}", 12)]])
        code, rows, _ = self.run_cli(["check", pdf, "--source", md])
        self.assertEqual((code, rows[0]["verdict"]), (0, "warn"), rows)

    def test_check_source_consumes_split_long_token_occurrences(self):
        self.require_tool("node")
        md = self.write_md("long-token.md", "abcdefghijkl abcdefghijkl\n")
        pdf = self.make_pdf("one-long-token.pdf", [[(50, 100, "abcdefghijkl", 12)]])
        code, rows, _ = self.run_cli(["check", pdf, "--source", md])
        self.assertEqual((code, rows[0]["verdict"]), (0, "warn"), rows)
        self.assertTrue(rows[0]["data"]["findings"], rows)

    def test_check_source_missing_words_are_judged_in_order_not_greedily(self):
        self.require_tool("node")
        md = self.write_md("order.md", "common one two three four\n\nfive six seven eight\n")
        pdf = self.make_pdf("order.pdf", [[(50, 100, "one two three four five six seven eight", 11),
                                           (50, 700, "common", 11)]])
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual((code, record["verdict"]), (0, "ok"), record)

    def test_check_source_consecutive_missing_run_warns_scattered_does_not(self):
        self.require_tool("node")
        words = "alpha bravo charlie delta echo foxtrot golf hotel india juliet kilo lima".split()
        md = self.write_md("runs.md", " ".join(words) + "\n")
        scattered = [w for w in words if w not in ("bravo", "foxtrot", "juliet")]
        pdf = self.make_pdf("scattered.pdf", [[(50, 100, " ".join(scattered), 11)]])
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual(code, 0)
        self.assertNotIn("MISSING_WORD_RUN", [w["code"] for w in record["data"]["warnings"]])
        consecutive = [w for w in words if w not in ("charlie", "delta", "echo")]
        pdf = self.make_pdf("consecutive.pdf", [[(50, 100, " ".join(consecutive), 11)]])
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual(code, 0)
        self.assertIn("MISSING_WORD_RUN", [w["code"] for w in record["data"]["warnings"]])

    def test_check_source_raw_tex_fails_outside_code_and_is_excused_in_code(self):
        self.require_tool("node")
        pdf = self.make_pdf("tex.pdf", [[(50, 100, "The ratio is \\frac{1}{2} here", 11)]])
        code, rows, stderr = self.run_cli(["check", pdf, "--source", self.write_md("tex-prose.md", "The ratio is half here\n")])
        self.assertEqual((code, rows[0]["verdict"]), (0, "warn"))
        self.assertIn("RAW_TEX", [finding["code"] for finding in rows[0]["data"]["findings"]])
        self.assertIn("WARN pdf check RAW_TEX", stderr)
        code, record, _ = self.source_check(pdf, self.write_md("tex-code.md", "The ratio is `\\frac{1}{2}` here\n"))
        self.assertEqual((code, record["verdict"]), (0, "ok"), record)

    def test_check_source_ordinary_inline_syntax_is_not_unchecked(self):
        self.require_tool("node")
        pdf = self.make_pdf("syntax.pdf", [[(50, 100, "Only a link in this paragraph", 11), (50, 130, "see code here", 11)]])
        md = self.write_md("syntax.md", "[Only a link in this paragraph](https://example.org)\n\n![](pic.png)\n\nsee `code` here\n")
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual(code, 0, record)
        self.assertEqual([w["code"] for w in record["data"]["warnings"] if w["code"] != "IMAGE_COUNT_CHANGED"], [])

    def test_check_source_treats_inline_html_break_as_whitespace(self):
        self.require_tool("node")
        pdf = self.make_pdf("html-break.pdf", [[(50, 100, "one two", 11)]])
        md = self.write_md("html-break.md", "one<br>two\n")
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual(code, 0, record)
        self.assertEqual(record["verdict"], "ok")

    def test_check_source_geometry_warnings(self):
        self.require_tool("node")
        md = self.write_md("geometry.md", "# Orphan heading\n\nBody text starts on the next page.\n\nLow line of text\n")
        pdf = self.make_pdf("geometry.pdf", [
            [(57, 790, "Orphan heading", 16)],
            [(57, 100, "Body text starts on the next page.", 11), (57, 815, "Low line of text", 11),
             (57, 300, "tiny one", 5), (57, 320, "tiny two", 5)]])
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual(code, 0, record)
        codes = [(w["code"], w.get("page")) for w in record["data"]["warnings"]]
        self.assertIn(("HEADING_ORPHAN", 1), codes)
        self.assertIn(("BLOCK_CUT", 2), codes)
        self.assertEqual(codes.count(("SMALL_TEXT", 2)), 1)

    def test_check_source_image_cut_by_the_bottom_margin_warns(self):
        self.require_tool("node")
        md = self.write_md("cut.md", "Some words before the picture.\n")
        path = self.root / "cut.pdf"
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((57, 100), "Some words before the picture.", fontsize=11)
        page.insert_image(pymupdf.Rect(57, 700, 200, 835), filename=str(self.fixtures["image"]))
        doc.save(path)
        doc.close()
        code, record, _ = self.source_check(path, md)
        self.assertEqual(code, 0, record)
        self.assertIn(("BLOCK_CUT", 1), [(w["code"], w.get("page")) for w in record["data"]["warnings"]])

    def test_check_source_large_document_finishes_quickly(self):
        self.require_tool("node")
        import random
        rng = random.Random(7)
        vocabulary = [f"term{i}" for i in range(60)]
        paragraphs = [" ".join(rng.choice(vocabulary) for _ in range(80)) for _ in range(60)]
        md = self.write_md("large.md", "\n\n".join(paragraphs) + "\n")
        text = self.write_md("large.txt", "\n".join(paragraphs) + "\n")
        pdf = self.root / "large.pdf"
        self.assertEqual(self.run_cli(["txt2pdf", text, "-o", pdf])[0], 0)
        started = time.monotonic()
        code, record, _ = self.source_check(pdf, md)
        self.assertEqual((code, record["verdict"]), (0, "ok"), record)
        self.assertLess(time.monotonic() - started, 60)

    def test_converters_name_unreadable_inputs(self):
        for command_name, extension in (("md2pdf", ".md"), ("txt2pdf", ".txt"), ("office2pdf", ".docx"), ("toword", ".pdf")):
            code, rows, _ = self.run_cli([command_name, self.root / f"absent{extension}"])
            self.assertEqual((code, rows[0]["code"]), (1, "NOT_FOUND"), command_name)
    # ---- fix round 4 regressions ------------------------------------------------------------------

    def test_docx_word_count_counts_a_merged_cell_once(self):
        from docx import Document
        document = Document()
        table = document.add_table(rows=1, cols=10)
        merged = table.cell(0, 0).merge(table.cell(0, 9))
        merged.text = " ".join(f"word{i}" for i in range(20))
        self.assertEqual(p6_convert._docx_words(document), 20)  # row.cells repeats the merged cell ten times

    def test_heading_orphan_requires_a_long_prefix_or_exact_match(self):
        pdf = self.root / "short-heading.pdf"
        with pymupdf.open() as doc:
            first = doc.new_page()
            first.insert_text((50, 790), "in the", fontsize=11)
            second = doc.new_page()
            second.insert_text((50, 80), "Within the model", fontsize=16)
            second.insert_text((50, 110), "This body text follows the heading.", fontsize=11)
            doc.save(pdf)
        with pymupdf.open(pdf) as doc:
            result = p6_convert._compare_source(doc, [
                {"type": "paragraph", "text": "in the", "section": "doc"},
                {"type": "heading", "text": "Within the model", "section": "doc"},
                {"type": "paragraph", "text": "This body text follows the heading.", "section": "doc"},
            ], 0, pdf, pdf)
        self.assertNotIn("HEADING_ORPHAN", [warning["code"] for warning in result["data"]["warnings"]])

    def test_txt2pdf_closes_document_writer_when_story_placement_raises(self):
        source = self.write_md("writer.txt", "text")
        writer = mock.Mock()
        story = mock.Mock()
        story.place.side_effect = RuntimeError("place failed")
        def capture(src, args, context, writer_fn, *, command_name, verify):
            with self.assertRaisesRegex(RuntimeError, "place failed"):
                writer_fn(src, self.root / "stage.pdf", self.root)
        with mock.patch.object(p6_convert.pymupdf, "Story", return_value=story), \
             mock.patch.object(p6_convert.pymupdf, "DocumentWriter", return_value=writer), \
             mock.patch.object(p6_convert, "_publish_pdf", side_effect=capture):
            p6_convert.txt2pdf(mock.Mock(file=source, output=None, dry_run=False, mono=False), mock.Mock())
        writer.close.assert_called_once()

    def test_failed_page_note_reads_the_error_line_not_the_progress_lines(self):
        text = "(1/7) Page 1\n(2/7) Page 2\n[ERROR] Error when parsing page 7: boom\n"
        self.assertEqual(p6_convert._failed_page_note(text), " on page 7")
        self.assertEqual(p6_convert._failed_page_note("(1/7) Page 1\nnothing failed here"), "")

    def test_translated_tool_failure_keeps_the_captured_tool_output(self):
        run = subprocess.CompletedProcess(["tool", "arg"], 3, "captured-stdout", "captured-stderr-line")
        error = p6_convert._fail_tool(p6_convert.clierr.Fail("TOOL_FAILED", "tool exit 3", "see log", tool_run=run))
        self.assertIn("captured-stderr-line", error.detail)
        self.assertIn("captured-stdout", error.detail)
        self.assertIn("tool", error.detail)

    def test_signed_pdf_exports_warn_that_the_signature_is_not_retained(self):
        import pikepdf
        signed = self.root / "signed-export.pdf"
        with pikepdf.open(self.fixtures["pdf"]) as pdf:
            signature = pdf.make_indirect(pikepdf.Dictionary({"/FT": pikepdf.Name("/Sig"), "/V": pikepdf.Dictionary(
                {"/Type": pikepdf.Name("/Sig"), "/Contents": pikepdf.String("x")})}))
            pdf.Root.AcroForm = pikepdf.Dictionary({"/Fields": pikepdf.Array([signature])})
            pdf.save(signed)
        for command_name in ("toword", "images"):
            code, rows, _ = self.run_cli(["--dry-run", command_name, signed])
            self.assertEqual(code, 0, rows)
            self.assertIn("SIGNATURE_NOT_RETAINED", rows[0]["warnings"], command_name)

    def collision(self, command_name, source, name, extension):
        """Run twice into the same -o; the second publication is renamed and says so with a warning."""
        output = self.root / f"collide-{name}" / f"out{extension}"
        output.parent.mkdir()
        first = self.run_cli([command_name, source, "-o", output])
        self.assertEqual(first[0], 0, first[1])
        code, rows, err = self.run_cli([command_name, source, "-o", output])
        self.assertEqual(code, 0, (rows, err))
        self.assertIn("NAME_TAKEN", rows[0]["warnings"], rows)
        self.assertEqual(rows[0]["verdict"], "warn", rows)
        self.assertNotEqual(Path(rows[0]["path"]), output)
        self.assertIn("NAME_TAKEN", err)

    def test_txt2pdf_name_collision_is_a_warning(self):
        self.collision("txt2pdf", self.fixtures["text"], "txt", ".pdf")

    def test_office2pdf_name_collision_is_a_warning(self):
        self.require_tool("soffice")
        self.collision("office2pdf", self.fixtures["docx"], "office", ".pdf")

    def test_toword_name_collision_is_a_warning(self):
        self.require_tool("pdf2docx")
        self.collision("toword", self.fixtures["pdf"], "word", ".docx")

    def test_md2pdf_name_collision_is_a_warning(self):
        self.require_markdown_tools()
        self.collision("md2pdf", self.fixtures["markdown"], "md", ".pdf")

    def test_img2pdf_name_collision_is_a_warning(self):
        self.collision("img2pdf", self.fixtures["image"], "img", ".pdf")

    def test_images_name_collision_is_a_warning(self):
        output = self.root / "collide-images" / "out"
        output.parent.mkdir()
        target = str(output) + "\\"
        self.assertEqual(self.run_cli(["images", self.fixtures["pdf"], "-o", target])[0], 0)
        code, rows, err = self.run_cli(["images", self.fixtures["pdf"], "-o", target])
        self.assertEqual(code, 0, (rows, err))
        self.assertIn("NAME_TAKEN", rows[0]["warnings"])
        self.assertEqual(rows[0]["verdict"], "warn")


if __name__ == "__main__":
    unittest.main()
