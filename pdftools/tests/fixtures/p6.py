"""Phase 6 input fixtures, generated locally without toolkit commands."""
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFont
from docx import Document


def build_fixtures(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    pdf = folder / "phase6.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=595)
    page.insert_text((42, 50), "Phase six conversion fixture", fontsize=16)
    page.insert_textbox(pymupdf.Rect(42, 80, 380, 240),
                        "Climate observation word sequence remains intact across a conversion. " * 8, fontsize=11)
    page.draw_rect(pymupdf.Rect(50, 280, 320, 500), color=(.2, .4, .8), width=2)
    doc.save(pdf)
    doc.close()

    text = folder / "notes.txt"
    text.write_text("A & B < C > D\nsecond line", encoding="utf-8-sig")
    markdown = folder / "notes.md"
    markdown.write_text("# Climate notes\n\nA clear **observation** with $x+1$ and a display equation.\n\n$$E=mc^2$$\n\nThe following paragraph contains enough words to verify that the rendered source is present in the PDF output and in the ordered comparison sequence.\n", encoding="utf-8")

    image = folder / "chart.png"
    im = Image.new("RGBA", (400, 250), (255, 255, 255, 255))
    draw = ImageDraw.Draw(im)
    draw.rectangle((10, 10, 389, 239), outline=(20, 70, 150, 255), width=5)
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 34)
    draw.text((30, 80), "Weather 2026", fill=(0, 0, 0, 255), font=font)
    im.save(image, dpi=(300, 300)); im.close()
    with pymupdf.open(pdf) as doc:
        doc[0].insert_image(pymupdf.Rect(50, 280, 320, 500), filename=str(image))
        replacement = pdf.with_suffix(".new.pdf")
        doc.save(replacement)
    replacement.replace(pdf)
    changed = folder / "chart-changed.png"
    im = Image.new("RGB", (400, 250), "white")
    ImageDraw.Draw(im).rectangle((6, 5, 394, 244), fill="red")
    im.save(changed); im.close()
    svg = folder / "vector.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100"><text x="10" y="40">Vector</text></svg>', encoding="utf-8")
    eps = folder / "ghostscript.eps"
    eps.write_text("%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 200 100\n1 0 0 setrgbcolor 10 10 180 80 rectfill\nshowpage\n", encoding="ascii")

    docx = folder / "office.docx"
    word = Document()
    word.add_heading("Office conversion", 0)
    word.add_paragraph("A document with a climate value and a local hyperlink target.")
    word.save(docx)
    return {"pdf": pdf, "text": text, "markdown": markdown, "image": image,
            "changed_image": changed, "svg": svg, "eps": eps, "docx": docx}


# --- converter fixtures (office2pdf, toword, md2pdf, txt2pdf) -----------------------------------

import zipfile
from xml.sax.saxutils import escape as _xml_escape

PASSWORD = 'pw "q" \\ ü中 %^&'  # quotes, backslash, Unicode, cmd metacharacters

_ODF_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
           'xmlns:script="urn:oasis:names:tc:opendocument:xmlns:script:1.0" '
           'xmlns:dom="http://www.w3.org/2001/xml-events" xmlns:xlink="http://www.w3.org/1999/xlink" xmlns:ooo="http://openoffice.org/2004/office"')
_LIBRARIES = ('<?xml version="1.0" encoding="UTF-8"?><library:libraries xmlns:library="http://openoffice.org/2000/library" '
              'xmlns:xlink="http://www.w3.org/1999/xlink"><library:library library:name="Standard" library:link="false"/>'
              '</library:libraries>')
_LIBRARY = ('<?xml version="1.0" encoding="UTF-8"?><library:library xmlns:library="http://openoffice.org/2000/library" '
            'library:name="Standard" library:readonly="false" library:passwordprotected="false">'
            '<library:element library:name="Module1"/></library:library>')


def _odt(path, body, scripts="", extra=()):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text")
        entries = "".join(f'<manifest:file-entry manifest:full-path="{n}" manifest:media-type="{t}"/>' for n, t, _ in extra)
        z.writestr("META-INF/manifest.xml",
                   '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">'
                   '<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>'
                   f'<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>{entries}</manifest:manifest>')
        z.writestr("content.xml", f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {_ODF_NS} '
                                  f'office:version="1.3">{scripts}<office:body><office:text>{body}</office:text></office:body></office:document-content>')
        for name, _, data in extra:
            z.writestr(name, data)


def build_office_tripwires(folder):
    """An .odt whose open event runs a Basic macro writing `marker`, and an .odt with a linked section
    whose source file holds FRESH text while the cached text is STALE."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    marker = folder / "macro-ran.txt"
    events = ('<office:scripts><office:script script:language="ooo:Basic"><ooo:libraries><ooo:library-embedded ooo:name="Standard"/>'
              '</ooo:libraries></office:script><office:event-listeners><script:event-listener script:language="ooo:script" '
              'script:event-name="dom:load" xlink:href="vnd.sun.star.script:Standard.Module1.Run?language=Basic&amp;location=document" '
              'xlink:type="simple"/></office:event-listeners></office:scripts>')
    module = ('<?xml version="1.0" encoding="UTF-8"?><!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">'
              '<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">'
              f'Sub Run\nDim n As Integer\nn = FreeFile\nOpen "{_xml_escape(str(marker).replace(chr(92), "/"))}" For Output As #n\n'
              'Print #n, "ran"\nClose #n\nEnd Sub\n</script:module>')
    _odt(folder / "macro.odt", "<text:p>Macro tripwire document</text:p>", events,
         [("Basic/", "", ""), ("Basic/Standard/", "", ""), ("Basic/script-lc.xml", "text/xml", _LIBRARIES),
          ("Basic/Standard/script-lb.xml", "text/xml", _LIBRARY), ("Basic/Standard/Module1.xml", "text/xml", module)])
    _odt(folder / "fresh.odt", "<text:p>FRESH-LINKED-TEXT</text:p>")
    uri = (folder / "fresh.odt").resolve().as_uri()
    _odt(folder / "link.odt", '<text:p>Before link</text:p><text:section text:name="Link1"><text:section-source '
         f'xlink:href="{uri}#" xlink:type="simple" text:filter-name="writer8"/><text:p>STALE-CACHED-TEXT</text:p></text:section>')
    return {"macro_odt": folder / "macro.odt", "marker": marker, "link_odt": folder / "link.odt"}


def build_encrypted_pdf(folder, password=PASSWORD):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "locked.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 500, 300),
                        "Encrypted conversion fixture paragraph with a number of ordinary words to convert. " * 4, fontsize=12)
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password + "-owner")
    doc.close()
    return path


def build_markdown_fixtures(folder):
    """Markdown cases: rich (code, links, image, list, table, math), BOM heading, equation sheet."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (60, 40), (200, 60, 60)).save(folder / "pic.png")
    rich = folder / "rich.md"
    rich.write_text(
        "# Rich document\n\n"
        "![a red square](pic.png)\n\n"
        "[Only a link in this paragraph](https://example.org/page)\n\n"
        "Inline code `\\frac{1}{2}` and `<img src=\"https://127.0.0.1:9/x.png\">` stay literal text, while $\\frac{1}{2}$ is math.\n\n"
        "- list item with a fenced block\n\n"
        "  ```\n  $$ not math $$ <b>literal</b>\n  ```\n\n"
        "Paragraph before the indented code.\n\n    indented code with $dollars$ and <i>tags</i>\n\n"
        "| name | value |\n|---|---|\n| alpha | `beta` |\n\n"
        "Costs $5 and $6 are plain dollars; \\$ is a literal sign.\n", encoding="utf-8")
    bom = folder / "bom.md"
    bom.write_bytes(b"\xef\xbb\xbf# Heading after BOM\n\nBody paragraph text for the byte order mark case.\n")
    terms = "+".join(f"a_{{{i}}}" for i in range(1, 36))
    equations = folder / "equations.md"
    equations.write_text(
        "# Equations\n\nBefore paragraph.\n\n"
        "$$\n\\begin{aligned}\n" + " \\\\\n".join(f"E_{n} &= {terms}" for n in (1, 2, 3)) + "\n\\end{aligned}\n$$\n\n"
        "AFTERWORD paragraph.\n", encoding="utf-8")
    return {"rich_md": rich, "bom_md": bom, "equations_md": equations, "pic": folder / "pic.png"}
